"""Полный исходный JSON v2.1 и экспорт таблиц, источников и дерева."""
from __future__ import annotations

import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

import pandas as pd
from openpyxl.styles import Alignment

from .bibliography import build_bibliography_text
from .data import build_source_school_summary
from .tables import (build_candidates_dataframe, build_classification_dataframes,
                     build_evidence_dataframe, build_information_dataframe, build_people_dataframe,
                     build_records_dataframe, build_sources_dataframe)
from .tree import SourceSchoolTree, build_source_school_tree_edges_dataframe


@dataclass(frozen=True)
class SourceSchoolExportBundle:
    json_bytes: bytes
    xlsx_bytes: bytes
    representatives_csv_bytes: bytes
    sources_csv_bytes: bytes
    evidence_csv_bytes: bytes
    bibliography_txt_bytes: bytes
    tree_html_bytes: bytes
    tree_png_bytes: bytes
    zip_bytes: bytes
    assignments_csv_bytes: bytes = b""
    assignment_bases_csv_bytes: bytes = b""


def _csv_bytes(dataframe: pd.DataFrame) -> bytes:
    return dataframe.to_csv(index=False).encode("utf-8-sig")


def _write_sheet(writer: pd.ExcelWriter, name: str, dataframe: pd.DataFrame) -> None:
    dataframe.to_excel(writer, sheet_name=name, index=False)
    sheet = writer.book[name]
    sheet.freeze_panes = "A2"
    if sheet.max_row and sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = min(max(len(str(c.value or "")) for c in cells) + 2, 60)
        for cell in cells:
            # Текст источника сохраняется как текст, без выполнения формул Excel.
            if cell.data_type == "f":
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def _tables(document: Mapping[str, Any], tree: SourceSchoolTree) -> dict[str, pd.DataFrame]:
    assignments, bases = build_classification_dataframes(document)
    metadata = {**document["метаданные"], **document["школа"], **build_source_school_summary(document)}
    return {
        "Общие сведения": build_records_dataframe([metadata]),
        "Представители": build_people_dataframe(document, display_only=False),
        "Кандидаты": build_candidates_dataframe(document),
        "Источники": build_sources_dataframe(document),
        "Подтверждения": build_evidence_dataframe(document),
        "Классификации": assignments,
        "Основания классификаций": bases,
        "Сведения": build_information_dataframe(document),
        "Связи": build_records_dataframe(document["связи"]),
        "Объекты": build_records_dataframe(document["объекты"]["прочие"]),
        "Состав": build_records_dataframe(document["состав"]),
        "Рёбра дерева": build_source_school_tree_edges_dataframe(tree),
        "Разночтения": build_records_dataframe(document["редактура"]["разночтения"]),
        "Редакторские решения": build_records_dataframe(document["редактура"]["журнал_решений"]),
    }


def _xlsx_bytes(tables: Mapping[str, pd.DataFrame]) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, dataframe in tables.items():
            _write_sheet(writer, name, dataframe)
    return output.getvalue()


def build_source_school_export_bundle(*, document: Mapping[str, Any], source_path: Path,
                                     tree: SourceSchoolTree, tree_html: str,
                                     tree_png: bytes) -> SourceSchoolExportBundle:
    """Экспортирует переданный текущий документ целиком.

    Аргумент ``source_path`` сохранён в API как контекст. Источником экспорта служит
    полный документ, поэтому правки и редакторские записи в памяти сохраняются в JSON.
    """
    slug = document["школа"]["идентификатор_школы"]
    json_bytes = (json.dumps(document, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    tables = _tables(document, tree)
    xlsx_bytes = _xlsx_bytes(tables)
    representatives_csv = _csv_bytes(tables["Представители"])
    sources_csv = _csv_bytes(tables["Источники"])
    evidence_csv = _csv_bytes(tables["Подтверждения"])
    assignments_csv = _csv_bytes(tables["Классификации"])
    bases_csv = _csv_bytes(tables["Основания классификаций"])
    bibliography = build_bibliography_text(document).encode("utf-8")
    html = tree_html.encode("utf-8")
    files = {
        f"{slug}.json": json_bytes, f"{slug}.данные.xlsx": xlsx_bytes,
        f"{slug}.представители.csv": representatives_csv, f"{slug}.источники.csv": sources_csv,
        f"{slug}.подтверждения.csv": evidence_csv,
        f"{slug}.классификационные_назначения.csv": assignments_csv,
        f"{slug}.основания_классификации.csv": bases_csv,
        f"{slug}.кандидаты.csv": _csv_bytes(tables["Кандидаты"]),
        f"{slug}.сведения.csv": _csv_bytes(tables["Сведения"]),
        f"{slug}.связи.csv": _csv_bytes(tables["Связи"]),
        f"{slug}.объекты.csv": _csv_bytes(tables["Объекты"]),
        f"{slug}.рёбра_дерева.csv": _csv_bytes(tables["Рёбра дерева"]),
        f"{slug}.список_источников.txt": bibliography,
        f"{slug}.интерактивное_дерево.html": html, f"{slug}.дерево.png": tree_png,
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return SourceSchoolExportBundle(json_bytes, xlsx_bytes, representatives_csv, sources_csv,
                                    evidence_csv, bibliography, html, tree_png, output.getvalue(),
                                    assignments_csv, bases_csv)


def build_filtered_people_csv(dataframe: pd.DataFrame) -> bytes:
    return _csv_bytes(dataframe.copy())


def build_filtered_people_xlsx(dataframe: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _write_sheet(writer, "Представители", dataframe.copy())
    return output.getvalue()
